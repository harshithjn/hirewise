import io
from pathlib import Path

import pdfplumber
from docx import Document

from errors import AppError

SUPPORTED = {".pdf", ".docx", ".txt"}
MAX_CHARS = 30_000


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED:
        raise AppError(f"Unsupported file type '{ext or 'unknown'}'. Use PDF, DOCX or TXT.", code="unsupported_file")
    try:
        if ext == ".pdf":
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                text = "\n".join((p.extract_text() or "") for p in pdf.pages)
        elif ext == ".docx":
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for t in doc.tables:
                parts += [" | ".join(c.text for c in row.cells) for row in t.rows]
            text = "\n".join(parts)
        else:
            try:
                text = data.decode("utf-8")
            except UnicodeDecodeError:
                text = data.decode("latin-1")
    except AppError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise AppError(f"Could not read the file ({ext[1:].upper()} may be corrupt or encrypted).", code="unreadable_file") from exc
    return text.strip()[:MAX_CHARS]
